from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock

import pytest

import file_processor
import ocr_runtime
import status_registry
from file_processor import FileProcessor
from status_registry import RuntimeStatus


@pytest.fixture(autouse=True)
def isolated_status(tmp_path, monkeypatch):
    registry = RuntimeStatus(tmp_path / "status.json", session_id="phase4", pid=404)
    monkeypatch.setattr(status_registry, "_registry_instance", registry)
    ocr_runtime.reset_tesseract_state()
    yield registry
    ocr_runtime.reset_tesseract_state()


def _processor():
    return FileProcessor()


def _write_text_pdf(path: Path, text: str):
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 40 200 Td ({safe}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        f"<< /Length {len(content)} >>\nstream\n".encode("ascii") + content + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    path.write_bytes(pdf)


def test_tesseract_found_via_explicit_path(tmp_path, monkeypatch):
    executable = tmp_path / "configured-tesseract.exe"
    executable.write_bytes(b"stub")
    monkeypatch.setenv("TESSERACT_PATH", str(executable))
    monkeypatch.setattr(ocr_runtime.shutil, "which", lambda _: None)
    result = ocr_runtime.discover_tesseract()
    assert result.path == str(executable.resolve())
    assert result.method == "TESSERACT_PATH"


def test_tesseract_found_via_path(tmp_path, monkeypatch):
    executable = tmp_path / "path-tesseract.exe"
    executable.write_bytes(b"stub")
    monkeypatch.delenv("TESSERACT_PATH", raising=False)
    monkeypatch.setattr(ocr_runtime.shutil, "which", lambda _: str(executable))
    result = ocr_runtime.discover_tesseract()
    assert result.path == str(executable.resolve())
    assert result.method == "PATH"


def test_tesseract_missing(monkeypatch):
    monkeypatch.setattr(ocr_runtime, "_candidate_paths", lambda: [])
    result = ocr_runtime.discover_tesseract()
    assert result.state == "UNAVAILABLE"
    assert result.path is None


def test_tesseract_probe_failure(monkeypatch):
    monkeypatch.setattr(ocr_runtime, "discover_tesseract", lambda: ocr_runtime.TesseractStatus(
        "DISCOVERED", "fake-tesseract", "test", "found"))
    monkeypatch.setattr(ocr_runtime.subprocess, "run", lambda *a, **k: SimpleNamespace(
        returncode=7, stdout="", stderr="bad"))
    result = ocr_runtime.probe_tesseract(force=True)
    assert result.state == "BROKEN"
    assert "code 7" in result.detail


def test_ocr_successful_extraction(monkeypatch, isolated_status):
    fake = SimpleNamespace(
        pytesseract=SimpleNamespace(tesseract_cmd=None),
        image_to_string=Mock(return_value="JARVIS OCR OK 2026\n"),
    )
    monkeypatch.setattr(ocr_runtime, "pytesseract", fake)
    monkeypatch.setattr(ocr_runtime, "probe_tesseract", lambda **k: ocr_runtime.TesseractStatus(
        "AVAILABLE", "fake-tesseract", "test", "ready"))
    image = object()
    text = ocr_runtime.extract_image_text(image, timeout=2)
    assert text == "JARVIS OCR OK 2026"
    fake.image_to_string.assert_called_once_with(image, timeout=2)
    assert isolated_status.get_capability("OCR")["evidence"] == "LIVE"


def test_ocr_timeout_is_safe(monkeypatch, isolated_status):
    fake = SimpleNamespace(
        pytesseract=SimpleNamespace(tesseract_cmd=None),
        image_to_string=Mock(side_effect=RuntimeError("process timeout")),
    )
    monkeypatch.setattr(ocr_runtime, "pytesseract", fake)
    monkeypatch.setattr(ocr_runtime, "probe_tesseract", lambda **k: ocr_runtime.TesseractStatus(
        "AVAILABLE", "fake-tesseract", "test", "ready"))
    with pytest.raises(ocr_runtime.OCRTimeoutError):
        ocr_runtime.extract_image_text(object(), timeout=0.1)
    assert isolated_status.get_capability("OCR")["evidence"] == "BROKEN"


def test_empty_ocr_is_not_meaningful_success(monkeypatch, isolated_status):
    fake = SimpleNamespace(
        pytesseract=SimpleNamespace(tesseract_cmd=None),
        image_to_string=Mock(return_value="   \n"),
    )
    monkeypatch.setattr(ocr_runtime, "pytesseract", fake)
    monkeypatch.setattr(ocr_runtime, "probe_tesseract", lambda **k: ocr_runtime.TesseractStatus(
        "AVAILABLE", "fake-tesseract", "test", "ready"))
    with pytest.raises(ocr_runtime.OCREmptyResultError):
        ocr_runtime.extract_image_text(object())
    assert isolated_status.get_capability("OCR")["evidence"] == "PROBED"


def test_unavailable_ocr_does_not_block_text_processing(tmp_path, monkeypatch):
    monkeypatch.setattr(ocr_runtime, "pytesseract", None)
    monkeypatch.setattr(ocr_runtime, "_candidate_paths", lambda: [])
    status = ocr_runtime.probe_tesseract(force=True)
    assert status.state == "UNAVAILABLE"
    source = tmp_path / "still-works.txt"
    source.write_text("TEXT WITHOUT OCR", encoding="utf-8")
    result = _processor().process_file(str(source), "extract")
    assert result.success is True
    assert result.result["text"] == "TEXT WITHOUT OCR"


def test_file_processor_ocr_promotes_image_capability(tmp_path, monkeypatch, isolated_status):
    from PIL import Image
    from status_registry import EvidenceLevel
    source = tmp_path / "ocr.png"
    image = Image.new("RGB", (40, 20), "white")
    image.save(source)
    image.close()
    isolated_status.set_evidence("TESSERACT_OCR", EvidenceLevel.LIVE, "mock OCR succeeded")
    monkeypatch.setattr(file_processor, "extract_image_text", lambda image, timeout=None: "OCR FILE OK")
    result = _processor().process_file(str(source), "ocr")
    assert result.success is True
    assert result.result == "OCR FILE OK"
    assert isolated_status.get_status("FILE_IMAGE_PARSER")["evidence"] == "LIVE"
    assert isolated_status.get_capability("FILE_IMAGE_OCR")["evidence"] == "LIVE"


@pytest.mark.parametrize("name,content", [
    ("notes.txt", "Unicode ✓ नमस्ते JARVIS"),
    ("sample.py", "def phase_four():\n    return 'CODE_OK'\n"),
])
def test_plain_text_and_code_extraction(tmp_path, name, content):
    source = tmp_path / name
    source.write_text(content, encoding="utf-8")
    result = _processor().process_file(str(source), "extract")
    assert result.success is True
    assert content.splitlines()[0] in result.result["text"]


def test_pdf_embedded_text_extraction(tmp_path):
    source = tmp_path / "text.pdf"
    _write_text_pdf(source, "PDF MARKER 2026")
    result = _processor().process_file(str(source), "extract")
    assert result.success is True
    assert "PDF MARKER 2026" in result.result["text"]


def test_pdf_without_extractable_text_is_truthful(tmp_path):
    from PyPDF2 import PdfWriter
    source = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with source.open("wb") as handle:
        writer.write(handle)
    result = _processor().process_file(str(source), "extract")
    assert result.success is False
    assert "no extractable text" in result.error
    assert "OCR is not implemented" in result.error


def test_docx_extraction_preserves_paragraph_order(tmp_path):
    import docx
    source = tmp_path / "document.docx"
    document = docx.Document()
    document.add_paragraph("FIRST PARAGRAPH")
    document.add_paragraph("SECOND PARAGRAPH")
    document.save(source)
    result = _processor().process_file(str(source), "extract")
    assert result.success is True
    text = result.result["text"]
    assert text.index("FIRST PARAGRAPH") < text.index("SECOND PARAGRAPH")


def test_xlsx_extraction_and_sheet_names(tmp_path):
    import openpyxl
    source = tmp_path / "workbook.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Phase4"
    sheet.append(["Name", "Value"])
    sheet.append(["JARVIS", 2026])
    workbook.save(source)
    workbook.close()
    result = _processor().process_file(str(source), "extract")
    assert result.success is True
    assert "--- Phase4 ---" in result.result["text"]
    assert "JARVIS\t2026" in result.result["text"]


def test_missing_optional_parser_is_not_success(tmp_path, monkeypatch):
    source = tmp_path / "input.pdf"
    source.write_bytes(b"%PDF-invalid")
    monkeypatch.setattr(file_processor, "HAS_PYPDF2", False)
    result = _processor().process_file(str(source), "extract")
    assert result.success is False
    assert "requires PyPDF2" in result.error


def test_corrupt_supported_file_fails_safely(tmp_path):
    source = tmp_path / "corrupt.docx"
    source.write_bytes(b"not a zip document")
    result = _processor().process_file(str(source), "extract")
    assert result.success is False
    assert "DOCX extraction failed" in result.error


def test_unsupported_file_type_fails_safely(tmp_path):
    source = tmp_path / "unknown.bin"
    source.write_bytes(b"binary")
    result = _processor().process_file(str(source), "extract")
    assert result.success is False
    assert "not allowed" in result.error


def test_source_file_is_releasable_after_repeated_processing(tmp_path):
    source = tmp_path / "release.xlsx"
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.active.append(["RELEASE", "OK"])
    workbook.save(source)
    workbook.close()
    processor = _processor()
    for _ in range(3):
        assert processor.process_file(str(source), "extract").success
    renamed = source.with_name("renamed.xlsx")
    source.rename(renamed)
    renamed.unlink()
    assert not renamed.exists()


def test_xlsx_cleanup_occurs_after_parser_exception(tmp_path, monkeypatch):
    source = tmp_path / "exception.xlsx"
    source.write_bytes(b"placeholder")
    worksheet = Mock()
    worksheet.iter_rows.side_effect = RuntimeError("parser exploded")
    workbook = MagicMock()
    workbook.sheetnames = ["Sheet"]
    workbook.__getitem__.return_value = worksheet
    monkeypatch.setattr(file_processor.openpyxl, "load_workbook", Mock(return_value=workbook))
    result = _processor().process_file(str(source), "extract")
    assert result.success is False
    assert workbook.close.call_count == 2  # metadata pass and extraction pass


def test_capability_truth_moves_from_configured_to_live(tmp_path, isolated_status):
    availability = file_processor.publish_parser_availability()
    assert all(info["available"] for info in availability.values())
    assert isolated_status.get_capability("FILE_TEXT")["evidence"] == "CONFIGURED"
    source = tmp_path / "truth.txt"
    source.write_text("TRUTHFUL EXTRACTION", encoding="utf-8")
    assert _processor().process_file(str(source), "extract").success
    assert isolated_status.get_status("FILE_TEXT_PARSER")["evidence"] == "LIVE"
    assert isolated_status.get_capability("FILE_TEXT")["evidence"] == "LIVE"
